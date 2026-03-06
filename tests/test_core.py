"""Unit tests for core matching and data extraction logic."""

import os
import re
import tempfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string


# ── Helpers extracted from beaver2.py for testability ────────────────────────


def extract_columns(ws, columns: list[str], start_row: int) -> list[list[str]]:
    """Extract specified columns from a worksheet starting at start_row.

    Returns a list of rows, each row is a list of string values.
    """
    column_indices = [column_index_from_string(col.strip()) - 1 for col in columns]
    result = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        current_values = [row[idx] if idx < len(row) else None for idx in column_indices]
        if all(v is None for v in current_values):
            continue
        formatted = [str(v) if v is not None else "" for v in current_values]
        result.append(formatted)
    return result


def find_matching_codes(numbers: list[str], elements: list[str]) -> list[str]:
    """Find elements whose codes match the given number patterns.

    Reproduces the finder logic from beaver2.py.
    """
    found = []
    for number in numbers:
        pattern = rf"^{re.escape(number)}(\.|$)"
        parts = number.split(".")
        last_part = parts[-1]

        if len(last_part) == 1 or (last_part.endswith("0") and len(last_part) > 1):
            base_number_prefix = ".".join(parts[:-1] + [last_part[:-1]]) if last_part.endswith("0") else number

            search_prefix = f"{base_number_prefix}"
            for element in elements:
                if not element:
                    continue
                try:
                    left_part = element.split()[0]
                    second_part = element.split()[-1]
                except IndexError:
                    continue
                if any(second_part.startswith(f"{search_prefix}{i}") for i in range(10)):
                    found.append(left_part)

        for element in elements:
            if not element:
                continue
            try:
                left_part = element.split()[0]
                second_part = element.split()[-1]
            except IndexError:
                continue
            if re.match(pattern, second_part):
                found.append(left_part)
    return found


def highlight_rows(
    ws,
    elements: list[str],
    search_col: str,
    color_hex: str,
    excluded_code: str | None = None,
):
    """Highlight rows where search_col value is in elements.

    If excluded_code is given, skip rows containing that code.
    """
    fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
    elements_set = set(elements)
    col_idx = column_index_from_string(search_col) - 1

    for row in ws.iter_rows():
        if len(row) <= col_idx:
            continue
        cell = row[col_idx]
        cell_value = str(cell.value).strip() if cell.value else ""
        if cell_value in elements_set:
            if excluded_code:
                has_forbidden = any(str(c.value).strip() == excluded_code for c in row)
                if has_forbidden:
                    continue
            for c in row:
                c.fill = fill


# ── Fixtures ────────────────────────────────────────────────────────────────


@pytest.fixture
def sample_workbook():
    """Create a workbook with sample procurement data."""
    wb = Workbook()
    ws = wb.active
    # Header rows (1-8)
    for i in range(1, 9):
        ws.cell(row=i, column=1, value=f"Header {i}")

    # Data rows starting from row 9
    data = [
        ("001", "Стулья офисные", "26.20.11", "Мебель", "100"),
        ("002", "Серверы Dell", "26.20.40", "IT-оборудование", "500"),
        ("003", "Бумага А4", "17.12.14.130", "Канцелярия", "50"),
        ("004", "Мониторы LG", "26.20.11.110", "IT-оборудование", "200"),
        ("005", "Картриджи", "XX.XX.XX.XXX", "Расходники", "75"),
    ]
    for row_idx, row_data in enumerate(data, start=9):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    return wb


@pytest.fixture
def temp_dir():
    """Provide a temporary directory for test files."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield tmpdir


# ── Tests: extract_columns ──────────────────────────────────────────────────


class TestExtractColumns:
    def test_extracts_correct_columns(self, sample_workbook):
        ws = sample_workbook.active
        result = extract_columns(ws, ["A", "C"], start_row=9)
        assert len(result) == 5
        assert result[0] == ["001", "26.20.11"]
        assert result[1] == ["002", "26.20.40"]

    def test_skips_empty_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="data")
        ws.cell(row=1, column=2, value="value")
        # Row 2 is empty
        ws.cell(row=3, column=1, value="more")
        ws.cell(row=3, column=2, value="stuff")
        result = extract_columns(ws, ["A", "B"], start_row=1)
        assert len(result) == 2

    def test_handles_none_values(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="data")
        ws.cell(row=1, column=2, value=None)
        result = extract_columns(ws, ["A", "B"], start_row=1)
        assert result[0] == ["data", ""]


# ── Tests: find_matching_codes ──────────────────────────────────────────────


class TestFindMatchingCodes:
    def test_exact_match(self):
        elements = ["001\t26.20.11", "002\t26.20.40"]
        result = find_matching_codes(["26.20.11"], elements)
        assert "001" in result

    def test_hierarchical_match(self):
        elements = ["001\t26.20.11.110", "002\t26.20.11.120", "003\t26.20.40"]
        result = find_matching_codes(["26.20.11"], elements)
        assert "001" in result
        assert "002" in result

    def test_no_match(self):
        elements = ["001\t99.99.99"]
        result = find_matching_codes(["26.20.11"], elements)
        assert "001" not in result

    def test_empty_elements(self):
        result = find_matching_codes(["26.20.11"], [])
        assert result == []

    def test_empty_numbers(self):
        result = find_matching_codes([], ["001\t26.20.11"])
        assert result == []

    def test_code_ending_in_zero(self):
        elements = ["001\t26.20.10", "002\t26.20.11", "003\t26.20.19"]
        result = find_matching_codes(["26.20.10"], elements)
        assert "001" in result

    def test_handles_malformed_elements(self):
        elements = ["", "   ", "no_tab_here"]
        result = find_matching_codes(["26.20"], elements)
        # Should not raise, graceful handling
        assert isinstance(result, list)


# ── Tests: highlight_rows ───────────────────────────────────────────────────


class TestHighlightRows:
    def test_highlights_matching_rows(self, sample_workbook):
        ws = sample_workbook.active
        highlight_rows(ws, ["001", "003"], search_col="A", color_hex="44944A")
        # Row 9 (001) should be highlighted
        assert ws.cell(row=9, column=1).fill.start_color.rgb == "0044944A"
        # Row 11 (003) should be highlighted
        assert ws.cell(row=11, column=1).fill.start_color.rgb == "0044944A"
        # Row 10 (002) should NOT be highlighted
        assert ws.cell(row=10, column=1).fill.start_color.rgb != "0044944A"

    def test_excludes_forbidden_code(self, sample_workbook):
        ws = sample_workbook.active
        highlight_rows(
            ws,
            ["005"],
            search_col="A",
            color_hex="44944A",
            excluded_code="XX.XX.XX.XXX",
        )
        # Row 13 (005) contains excluded code, should NOT be highlighted
        assert ws.cell(row=13, column=1).fill.start_color.rgb != "0044944A"

    def test_no_elements_no_highlight(self, sample_workbook):
        ws = sample_workbook.active
        highlight_rows(ws, [], search_col="A", color_hex="FF0000")
        for row_idx in range(9, 14):
            assert ws.cell(row=row_idx, column=1).fill.start_color.rgb != "00FF0000"


# ── Tests: File I/O ─────────────────────────────────────────────────────────


class TestFileOperations:
    def test_write_and_read_found_elements(self, temp_dir):
        path = os.path.join(temp_dir, "found_elements.txt")
        elements = ["001", "002", "003"]
        with open(path, "w", encoding="utf-8") as f:
            for el in elements:
                f.write(f"{el}\n")
        with open(path, encoding="utf-8") as f:
            read_back = [line.strip() for line in f if line.strip()]
        assert read_back == elements

    def test_workbook_save_load(self, sample_workbook, temp_dir):
        path = os.path.join(temp_dir, "test_output.xlsx")
        sample_workbook.save(path)
        assert os.path.exists(path)

        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(row=9, column=1).value == "001"
